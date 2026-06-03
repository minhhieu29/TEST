"""
Techcombank Mobile FAQ Crawler — chạy LOCAL trên máy của bạn

Yêu cầu:
  pip install playwright openpyxl
  playwright install chromium

Cách chạy:
  python tcb_mobile_faq_crawler.py --visible --slow --output techcombank_mobile_faq.xlsx
  python tcb_mobile_faq_crawler.py --output techcombank_mobile_faq.xlsx

Output:
  techcombank_mobile_faq.json
  techcombank_mobile_faq.csv
  techcombank_mobile_faq.xlsx
"""

import argparse
import csv
import json
import time
from collections import Counter

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout


FAQ_URLS = {
    "Techcombank Mobile": "https://techcombank.com/khach-hang-ca-nhan/ngan-hang-truc-tuyen/ngan-hang-so/techcombank-mobile/cau-hoi-thuong-gap",
}


def accept_cookies_if_any(page):
    possible_buttons = [
        "text=Đồng ý",
        "text=Chấp nhận",
        "text=Accept",
        "text=Accept all",
        "text=I agree",
    ]

    for selector in possible_buttons:
        try:
            btn = page.locator(selector).first
            if btn.is_visible(timeout=1500):
                btn.click(timeout=3000)
                time.sleep(0.8)
                print("  [cookie] Đã đóng popup cookie")
                return
        except Exception:
            pass


def expand_all(page, max_rounds=30):
    for round_num in range(max_rounds):
        buttons = []

        selectors = [
            "text=Xem thêm",
            "button:has-text('Xem thêm')",
            "[role='button']:has-text('Xem thêm')",
            "div:has-text('Xem thêm')",
        ]

        for selector in selectors:
            try:
                locators = page.locator(selector).all()
                for btn in locators:
                    try:
                        if btn.is_visible():
                            buttons.append(btn)
                    except Exception:
                        pass
            except Exception:
                pass

        unique_buttons = []
        seen_boxes = set()

        for btn in buttons:
            try:
                box = btn.bounding_box()
                if not box:
                    continue

                key = (
                    round(box["x"]),
                    round(box["y"]),
                    round(box["width"]),
                    round(box["height"]),
                )

                if key not in seen_boxes:
                    seen_boxes.add(key)
                    unique_buttons.append(btn)
            except Exception:
                pass

        if not unique_buttons:
            break

        print(f"  [expand] round {round_num + 1}: {len(unique_buttons)} nút 'Xem thêm'")

        clicked = 0

        for btn in unique_buttons:
            try:
                btn.scroll_into_view_if_needed()
                page.wait_for_timeout(200)
                btn.click(timeout=3000)
                clicked += 1
                page.wait_for_timeout(300)
            except Exception:
                pass

        if clicked == 0:
            break

        page.wait_for_timeout(1000)


def debug_page(page):
    try:
        title = page.title()
    except Exception:
        title = ""

    try:
        h2_count = page.locator("h2").count()
    except Exception:
        h2_count = 0

    try:
        h3_count = page.locator("h3").count()
    except Exception:
        h3_count = 0

    try:
        body_text = page.locator("body").inner_text(timeout=5000)
        body_len = len(body_text)
    except Exception:
        body_len = 0

    print(f"  Title: {title[:80]}")
    print(f"  h2 count: {h2_count}")
    print(f"  h3 count: {h3_count}")
    print(f"  body length: {body_len}")


def get_groups_and_questions(page):
    data = page.evaluate(
        r"""
        () => {
            const results = [];
            let currentGroup = "";

            const main =
                document.querySelector("main") ||
                document.querySelector("[role='main']") ||
                document.body;

            function cleanText(s) {
                return (s || "")
                    .replace(/\u00a0/g, " ")
                    .replace(/\s+/g, " ")
                    .replace(/expand_more/gi, "")
                    .replace(/chevron_left/gi, "")
                    .replace(/chevron_right/gi, "")
                    .replace(/keyboard_arrow_down/gi, "")
                    .replace(/keyboard_arrow_up/gi, "")
                    .replace(/Image/gi, "")
                    .trim();
            }

            function isBadText(text) {
                if (!text) return true;

                const badExact = [
                    "Xem thêm",
                    "Thu gọn",
                    "Xem tất cả",
                    "Tìm kiếm",
                    "Liên hệ",
                    "Đăng nhập",
                    "Mở tài khoản",
                    "Cá nhân",
                    "Doanh nghiệp",
                    "English",
                    "Tiếng Việt",
                    "Loại hình dịch vụ",
                    "Thông tin liên hệ"
                ];

                if (badExact.includes(text)) return true;

                const badContains = [
                    "©",
                    "Copyright",
                    "Facebook",
                    "Youtube",
                    "LinkedIn",
                    "App Store",
                    "Google Play",
                    "Bạn có muốn tiếp tục hành trình",
                    "Các tìm kiếm gần đây của bạn",
                    "Liên kết hữu ích",
                    "Điều khoản sử dụng",
                    "Chính sách bảo mật"
                ];

                return badContains.some(x => text.includes(x));
            }

            function cleanAnswer(answer, question) {
                let a = cleanText(answer);

                a = a
                    .replace(question, "")
                    .replace(/^Xem thêm/i, "")
                    .replace(/^Thu gọn/i, "")
                    .replace(/^expand_more/i, "")
                    .trim();

                return a;
            }

            const headings = Array.from(main.querySelectorAll("h2, h3"));

            for (let i = 0; i < headings.length; i++) {
                const el = headings[i];
                const tag = el.tagName;
                const text = cleanText(el.innerText || el.textContent);

                if (!text || text.length < 3) continue;

                if (tag === "H2") {
                    if (
                        text.includes("Câu hỏi thường gặp") ||
                        text.includes("Liên kết") ||
                        text.includes("Bạn có muốn") ||
                        text.length > 200
                    ) {
                        continue;
                    }

                    currentGroup = text;
                    continue;
                }

                if (tag === "H3") {
                    const question = text;

                    if (isBadText(question)) continue;
                    if (question.length < 8) continue;
                    if (question.length > 500) continue;

                    const nextHeading = headings[i + 1];

                    let answer = "";

                    try {
                        const range = document.createRange();
                        range.setStartAfter(el);

                        if (nextHeading) {
                            range.setEndBefore(nextHeading);
                        } else {
                            range.setEndAfter(main);
                        }

                        answer = cleanAnswer(range.toString(), question);
                    } catch (e) {
                        answer = "";
                    }

                    if (!answer || answer.length < 5) {
                        let parent = el.parentElement;

                        for (let depth = 0; depth < 5 && parent; depth++) {
                            const parentText = cleanAnswer(parent.innerText, question);

                            if (
                                parentText &&
                                parentText.length > 5 &&
                                parentText.length < 3000 &&
                                !isBadText(parentText)
                            ) {
                                answer = parentText;
                                break;
                            }

                            parent = parent.parentElement;
                        }
                    }

                    if (!answer || answer.length < 5) {
                        let parts = [];
                        let node = el;

                        for (let step = 0; step < 20; step++) {
                            node = node.nextElementSibling;

                            if (!node) break;

                            if (node.tagName === "H2" || node.tagName === "H3") {
                                break;
                            }

                            const t = cleanAnswer(node.innerText || node.textContent, question);

                            if (
                                t &&
                                t.length > 5 &&
                                t.length < 2000 &&
                                !isBadText(t)
                            ) {
                                parts.push(t);
                            }
                        }

                        answer = parts.join("\n").trim();
                    }

                    if (
                        question &&
                        answer &&
                        answer.length > 5 &&
                        !isBadText(answer)
                    ) {
                        results.push({
                            group: currentGroup,
                            question: question,
                            answer: answer.substring(0, 3000)
                        });
                    }
                }
            }

            return results;
        }
        """
    )

    return data


def crawl_page(page, category, url):
    print(f"\n[{category}] {url}")

    try:
        page.goto(url, wait_until="domcontentloaded", timeout=60000)
    except PWTimeout:
        print("  TIMEOUT khi load trang!")
        return []

    try:
        page.wait_for_load_state("networkidle", timeout=15000)
    except PWTimeout:
        pass

    page.wait_for_timeout(3000)

    accept_cookies_if_any(page)

    debug_page(page)

    try:
        h3_count = page.locator("h3").count()
    except Exception:
        h3_count = 0

    if h3_count == 0:
        print("  Trang chưa render đủ H3, chờ thêm...")
        page.wait_for_timeout(5000)
        debug_page(page)

    expand_all(page)

    page.wait_for_timeout(1500)

    items = get_groups_and_questions(page)

    print(f"  → Lấy được {len(items)} câu hỏi")

    if len(items) > 0:
        print("  Câu đầu tiên:")
        print("  Q:", items[0]["question"])
        print("  A:", items[0]["answer"][:300])

    for item in items:
        item["category"] = category
        item["url"] = url

    return items


def save_json(data, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"✓ Saved JSON: {path}")


def save_csv(data, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        fieldnames = [
            "stt",
            "category",
            "group",
            "question",
            "answer",
            "url",
        ]

        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()

        for i, row in enumerate(data, 1):
            writer.writerow(
                {
                    "stt": i,
                    "category": row.get("category", ""),
                    "group": row.get("group", ""),
                    "question": row.get("question", ""),
                    "answer": row.get("answer", ""),
                    "url": row.get("url", ""),
                }
            )

    print(f"✓ Saved CSV: {path}")


def save_excel(data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "FAQ Techcombank Mobile"

    headers = [
        "STT",
        "Danh mục",
        "Nhóm",
        "Câu hỏi",
        "Câu trả lời",
        "URL",
    ]

    ws.append(headers)

    header_fill = PatternFill("solid", start_color="CC0000")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)

    border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    alt_fill = PatternFill("solid", start_color="FFF5F5")
    normal_font = Font(name="Arial", size=10)
    question_font = Font(bold=True, name="Arial", size=10)

    for i, row in enumerate(data, 1):
        ws.append(
            [
                i,
                row.get("category", ""),
                row.get("group", ""),
                row.get("question", ""),
                row.get("answer", ""),
                row.get("url", ""),
            ]
        )

        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=i + 1, column=col)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = question_font if col == 4 else normal_font

            if i % 2 == 0:
                cell.fill = alt_fill

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 55
    ws.column_dimensions["E"].width = 80
    ws.column_dimensions["F"].width = 70

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)

    print(f"✓ Saved Excel: {path}")


def deduplicate(data):
    seen = set()
    deduped = []

    for item in data:
        category = item.get("category", "").strip().lower()
        question = item.get("question", "").strip().lower()
        answer = item.get("answer", "").strip().lower()

        key = (category, question, answer[:100])

        if question and len(question) > 5 and key not in seen:
            seen.add(key)
            deduped.append(item)

    return deduped


def main():
    parser = argparse.ArgumentParser(description="Crawl FAQ Techcombank Mobile")

    parser.add_argument(
        "--output",
        default="techcombank_mobile_faq.xlsx",
        help="Tên file output. Ví dụ: techcombank_mobile_faq.xlsx",
    )

    parser.add_argument(
        "--visible",
        action="store_true",
        help="Hiện cửa sổ browser để debug",
    )

    parser.add_argument(
        "--slow",
        action="store_true",
        help="Chạy chậm hơn để debug",
    )

    args = parser.parse_args()

    headless = not args.visible
    all_data = []

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=headless,
            slow_mo=200 if args.slow else 0,
        )

        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
            viewport={"width": 1440, "height": 900},
            locale="vi-VN",
            timezone_id="Asia/Ho_Chi_Minh",
        )

        page = context.new_page()

        for category, url in FAQ_URLS.items():
            items = crawl_page(page, category, url)
            all_data.extend(items)

            time.sleep(1.5)

        browser.close()

    deduped = deduplicate(all_data)

    print("\n" + "=" * 60)
    print(f"Raw data: {len(all_data)} dòng")
    print(f"Sau dedup: {len(deduped)} dòng")
    print("=" * 60)

    base = args.output.rsplit(".", 1)[0]

    save_json(deduped, base + ".json")
    save_csv(deduped, base + ".csv")
    save_excel(deduped, base + ".xlsx")

    cat_count = Counter(d.get("category", "") for d in deduped)

    print("\nPhân bố theo danh mục:")
    for cat, cnt in sorted(cat_count.items()):
        print(f"  {cat}: {cnt} câu hỏi")

    print("\nHoàn tất.")


if __name__ == "__main__":
    main()